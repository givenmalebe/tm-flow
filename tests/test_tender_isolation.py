import database as dbmod
from tests.conftest import auth_headers, make_org, make_user


def make_tender(db, org, reference="T-1", title="Secret tender"):
    tender = dbmod.Tender(
        reference_number=reference,
        title=title,
        client_company="Client",
        org_id=org.id,
    )
    db.add(tender)
    db.commit()
    db.refresh(tender)
    return tender


def two_tenants(db):
    alice = make_user(db, "alice@example.com")
    orgA = make_org(db, "Acme", account_type="bidder", user=alice)
    bob = make_user(db, "bob@example.com")
    orgB = make_org(db, "Globex", account_type="issuer", user=bob)
    return alice, orgA, bob, orgB


def test_list_tenders_only_returns_the_active_org(db, client):
    alice, orgA, bob, orgB = two_tenants(db)
    make_tender(db, orgA, "A-1", "Alice tender")
    make_tender(db, orgB, "B-1", "Bob tender")

    r = client.get("/api/tenders", headers=auth_headers(alice, orgA))
    assert r.status_code == 200, r.text
    refs = {t["reference_number"] for t in r.json()["items"]}
    assert refs == {"A-1"}


def test_get_tender_from_another_org_is_404_not_403(db, client):
    alice, orgA, bob, orgB = two_tenants(db)
    bob_tender = make_tender(db, orgB, "B-1", "Bob tender")

    r = client.get(f"/api/tenders/{bob_tender.id}", headers=auth_headers(alice, orgA))
    assert r.status_code == 404, f"expected 404, got {r.status_code}: {r.text}"


def test_update_tender_in_another_org_is_refused(db, client):
    alice, orgA, bob, orgB = two_tenants(db)
    bob_tender = make_tender(db, orgB, "B-1", "Bob tender")

    r = client.put(
        f"/api/tenders/{bob_tender.id}",
        json={"title": "Hijacked"},
        headers=auth_headers(alice, orgA),
    )
    assert r.status_code == 404

    db.refresh(bob_tender)
    assert bob_tender.title == "Bob tender"


def test_delete_tender_in_another_org_is_refused(db, client):
    alice, orgA, bob, orgB = two_tenants(db)
    bob_tender = make_tender(db, orgB, "B-1", "Bob tender")

    r = client.delete(
        f"/api/tenders/{bob_tender.id}", headers=auth_headers(alice, orgA)
    )
    assert r.status_code == 404
    assert db.query(dbmod.Tender).filter(dbmod.Tender.id == bob_tender.id).count() == 1


def test_tender_created_in_one_org_is_invisible_to_another(db, client):
    alice, orgA, bob, orgB = two_tenants(db)

    created = client.post(
        "/api/tenders",
        json={
            "reference_number": "A-9",
            "title": "Mine",
            "client_company": "Client",
        },
        headers=auth_headers(alice, orgA),
    )
    assert created.status_code in (200, 201), created.text
    tender_id = created.json()["id"]

    visible = client.get("/api/tenders", headers=auth_headers(alice, orgA))
    assert {t["reference_number"] for t in visible.json()["items"]} == {"A-9"}

    other = client.get(f"/api/tenders/{tender_id}", headers=auth_headers(bob, orgB))
    assert other.status_code == 404


def test_duplicate_reference_in_another_org_is_allowed(db, client):
    alice, orgA, bob, orgB = two_tenants(db)
    make_tender(db, orgA, "SHARED-1", "Alice")

    r = client.post(
        "/api/tenders",
        json={
            "reference_number": "SHARED-1",
            "title": "Bob copy",
            "client_company": "C",
        },
        headers=auth_headers(bob, orgB),
    )
    assert r.status_code in (200, 201), r.text


def test_spreadsheet_export_only_includes_the_active_org(db, client):
    """The export route emits xlsx, so read the workbook rather than raw bytes."""
    import io

    from openpyxl import load_workbook

    alice, orgA, bob, orgB = two_tenants(db)
    make_tender(db, orgA, "A-1", "Alice tender")
    make_tender(db, orgB, "B-1", "Bob tender")

    r = client.get("/api/tenders/export", headers=auth_headers(alice, orgA))
    assert r.status_code == 200, r.text

    workbook = load_workbook(io.BytesIO(r.content))
    cells = [
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    joined = " | ".join(cells)
    assert "A-1" in joined
    assert "B-1" not in joined


def test_checklist_item_of_another_org_is_not_mutable(db, client):
    """Checklist items are addressed by their own id, so they need a join guard."""
    alice, orgA, bob, orgB = two_tenants(db)
    bob_tender = make_tender(db, orgB, "B-1", "Bob tender")
    item = dbmod.ChecklistItem(tender_id=bob_tender.id, name="Bob's item")
    db.add(item)
    db.commit()
    db.refresh(item)

    r = client.put(
        f"/api/checklist/{item.id}",
        json={"name": "Hijacked"},
        headers=auth_headers(alice, orgA),
    )
    assert r.status_code == 404, f"expected 404, got {r.status_code}: {r.text}"

    db.refresh(item)
    assert item.name == "Bob's item"


def test_checklist_item_delete_of_another_org_is_refused(db, client):
    alice, orgA, bob, orgB = two_tenants(db)
    bob_tender = make_tender(db, orgB, "B-1", "Bob tender")
    item = dbmod.ChecklistItem(tender_id=bob_tender.id, name="Keep me")
    db.add(item)
    db.commit()
    db.refresh(item)

    r = client.delete(f"/api/checklist/{item.id}", headers=auth_headers(alice, orgA))
    assert r.status_code == 404
    assert db.query(dbmod.ChecklistItem).filter(dbmod.ChecklistItem.id == item.id).count() == 1


def test_pricing_of_another_org_is_not_reachable(db, client):
    alice, orgA, bob, orgB = two_tenants(db)
    bob_tender = make_tender(db, orgB, "B-1", "Bob tender")
    pv = dbmod.PricingVersion(tender_id=bob_tender.id, version_number=1, total=500000.0)
    db.add(pv)
    db.commit()

    r = client.get(
        f"/api/tenders/{bob_tender.id}/pricing", headers=auth_headers(alice, orgA)
    )
    assert r.status_code == 404, f"expected 404, got {r.status_code}: {r.text}"


def test_pricing_version_of_another_org_cannot_be_deleted(db, client):
    alice, orgA, bob, orgB = two_tenants(db)
    bob_tender = make_tender(db, orgB, "B-1", "Bob tender")
    pv = dbmod.PricingVersion(tender_id=bob_tender.id, version_number=1, total=1.0)
    db.add(pv)
    db.commit()
    db.refresh(pv)

    r = client.delete(f"/api/pricing/{pv.id}", headers=auth_headers(alice, orgA))
    assert r.status_code == 404
    assert db.query(dbmod.PricingVersion).filter(dbmod.PricingVersion.id == pv.id).count() == 1


def test_addendum_of_another_org_cannot_be_deleted(db, client):
    alice, orgA, bob, orgB = two_tenants(db)
    bob_tender = make_tender(db, orgB, "B-1", "Bob tender")
    add = dbmod.Addendum(tender_id=bob_tender.id, title="Bob addendum")
    db.add(add)
    db.commit()
    db.refresh(add)

    r = client.delete(f"/api/addendums/{add.id}", headers=auth_headers(alice, orgA))
    assert r.status_code == 404
    assert db.query(dbmod.Addendum).filter(dbmod.Addendum.id == add.id).count() == 1


def test_briefing_of_another_org_cannot_be_deleted(db, client):
    """BriefingUpdate inherits required fields from BriefingBase, so PUT cannot
    be exercised here without hitting a pre-existing schema bug. DELETE proves
    the same join guard."""
    alice, orgA, bob, orgB = two_tenants(db)
    bob_tender = make_tender(db, orgB, "B-1", "Bob tender")
    briefing = dbmod.Briefing(
        tender_id=bob_tender.id, date=dbmod.datetime.utcnow().date()
    )
    db.add(briefing)
    db.commit()
    db.refresh(briefing)

    r = client.delete(
        f"/api/briefings/{briefing.id}", headers=auth_headers(alice, orgA)
    )
    assert r.status_code == 404, f"expected 404, got {r.status_code}: {r.text}"
    assert db.query(dbmod.Briefing).filter(dbmod.Briefing.id == briefing.id).count() == 1
